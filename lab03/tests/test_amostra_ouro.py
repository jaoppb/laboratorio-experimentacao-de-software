"""Testes para o módulo de amostragem-ouro e geração de planilhas (Issue #63)."""

from datetime import datetime
from pathlib import Path
import json
import shutil

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from analises.amostra_ouro import (
    DEFAULT_SEED,
    TAG_PLACEHOLDER,
    SorteioInconsistenteError,
    completar_sorteio_pendente,
    executar_sorteio_amostra_ouro,
    OPCOES_ENTREGA_REAL,
    OPCOES_RELEASE_CORRETIVA,
    OPCOES_TIPO_PROJETO,
    ReleaseItem,
    RepoSampleItem,
    construir_dados_amostra_ouro,
    exportar_planilhas,
    gerar_dataframe_rotulacao,
    salvar_planilha_excel,
    sortear_releases_repositorio,
    sortear_repositorios,
)


@pytest.fixture
def dummy_repos_df():
    return pd.DataFrame(
        [
            {
                "owner": f"owner{i}",
                "name": f"repo{i}",
                "full_name": f"owner{i}/repo{i}",
                "stars": 1000 + i * 10,
                "language": "Python" if i % 2 == 0 else "TypeScript",
                "releases_count_window": 10 + i,
            }
            for i in range(100)
        ]
    )


def test_sortear_repositorios_reprodutivel(tmp_path: Path, dummy_repos_df: pd.DataFrame):
    parquet_path = tmp_path / "repos.parquet"
    dummy_repos_df.to_parquet(parquet_path)

    sample1 = sortear_repositorios(parquet_path, n=60, seed=42)
    sample2 = sortear_repositorios(parquet_path, n=60, seed=42)
    sample_diff_seed = sortear_repositorios(parquet_path, n=60, seed=99)

    assert len(sample1) == 60
    assert len(sample2) == 60
    assert sample1["full_name"].tolist() == sample2["full_name"].tolist()
    assert sample1["full_name"].tolist() != sample_diff_seed["full_name"].tolist()


def test_sortear_releases_reprodutivel():
    releases = [
        ReleaseItem(
            tag_name=f"v1.{i}.0",
            published_at=f"2026-0{i+1}-15T12:00:00Z",
            release_url=f"https://github.com/org/repo/releases/tag/v1.{i}.0",
        )
        for i in range(10)
    ]

    s1 = sortear_releases_repositorio(releases, n=5, seed=123)
    s2 = sortear_releases_repositorio(releases, n=5, seed=123)
    s_diff = sortear_releases_repositorio(releases, n=5, seed=999)

    assert len(s1) == 5
    assert [r.tag_name for r in s1] == [r.tag_name for r in s2]
    assert [r.tag_name for r in s1] != [r.tag_name for r in s_diff]

    # Garante que as 5 sorteadas estão em ordem cronológica
    dates = [r.published_at for r in s1]
    assert dates == sorted(dates)


def test_sortear_releases_poucas_releases():
    releases = [
        ReleaseItem(
            tag_name=f"v1.{i}.0",
            published_at=f"2026-0{i+1}-01T00:00:00Z",
            release_url=f"https://github.com/org/repo/releases/tag/v1.{i}.0",
        )
        for i in range(3)
    ]
    sampled = sortear_releases_repositorio(releases, n=5, seed=42)
    assert len(sampled) == 3
    assert [r.tag_name for r in sampled] == ["v1.0.0", "v1.1.0", "v1.2.0"]


def test_dataframe_rotulacao_sem_colunas_heuristica():
    """Garante que NENHUMA coluna da heurística é vazada para a planilha de rotulação."""
    rel = ReleaseItem(
        tag_name="v1.0.0",
        published_at="2026-01-01T00:00:00Z",
        release_url="https://github.com/org/repo/releases/tag/v1.0.0",
    )
    item = RepoSampleItem(
        repo_idx=1,
        full_name="org/repo",
        owner="org",
        name="repo",
        repo_url="https://github.com/org/repo",
        stars=5000,
        language="Python",
        releases_count_window=10,
        releases=[rel],
    )

    df = gerar_dataframe_rotulacao([item])

    # Colunas proibidas da heurística
    colunas_proibidas = {
        "cfr_b",
        "heuristic",
        "heuristica",
        "is_corrective",
        "corrective",
        "commits_count",
        "has_fix",
        "lead_time_a_hours",
        "lead_time_b_hours",
    }

    colunas_atuais = set(df.columns)
    assert not (colunas_proibidas & colunas_atuais), f"Colunas da heurística encontradas: {colunas_proibidas & colunas_atuais}"

    # Colunas obrigatórias
    colunas_esperadas = {
        "repo_idx",
        "repo_name",
        "repo_url",
        "tipo_projeto",
        "entrega_real",
        "release_idx",
        "tag_name",
        "published_at",
        "release_url",
        "release_corretiva",
        "notas_justificativa",
    }
    assert colunas_esperadas.issubset(colunas_atuais)

    # Campos de preenchimento devem estar vazios
    assert df.loc[0, "tipo_projeto"] == ""
    assert df.loc[0, "entrega_real"] == ""
    assert df.loc[0, "release_corretiva"] == ""


def test_exportar_planilhas_csv_e_excel(tmp_path: Path):
    rel = ReleaseItem(
        tag_name="v1.0.0",
        published_at="2026-01-01T00:00:00Z",
        release_url="https://github.com/org/repo/releases/tag/v1.0.0",
    )
    item = RepoSampleItem(
        repo_idx=1,
        full_name="org/repo",
        owner="org",
        name="repo",
        repo_url="https://github.com/org/repo",
        stars=1000,
        language="Rust",
        releases_count_window=5,
        releases=[rel],
    )
    df = gerar_dataframe_rotulacao([item])

    avaliadores = [
        {"id": "marcela", "nome": "Marcela Campos"},
        {"id": "gabriel", "nome": "Gabriel Assis"},
        {"id": "joao", "nome": "João Pedro Peres"},
    ]

    arquivos = exportar_planilhas(df, tmp_path, avaliadores=avaliadores)

    # Verifica que foram gerados arquivos para cada avaliador + template
    assert "marcela" in arquivos
    assert "gabriel" in arquivos
    assert "joao" in arquivos
    assert "template" in arquivos

    for av_id in ["marcela", "gabriel", "joao", "template"]:
        csv_file, xlsx_file = arquivos[av_id]
        assert csv_file.is_file()
        assert xlsx_file.is_file()

        # Valida CSV
        df_csv = pd.read_csv(csv_file)
        assert len(df_csv) == 1
        assert df_csv.loc[0, "repo_name"] == "org/repo"

        # Valida Excel com openpyxl
        wb = openpyxl.load_workbook(xlsx_file)
        assert "Rotulacao_Amostra_Ouro" in wb.sheetnames
        assert "Instrucoes" in wb.sheetnames
        ws = wb["Rotulacao_Amostra_Ouro"]
        assert ws.cell(row=1, column=1).value == "repo_idx"
        assert ws.cell(row=2, column=2).value == "org/repo"


# --------------------------------------------------------------------------- #
# Conclusão do sorteio pendente (correção da Issue #63)
# --------------------------------------------------------------------------- #
NUM_REPOS = 8
JANELA = {"window_start_str": "2025-10-01", "window_end_str": "2026-09-30"}


def _gravar_releases(releases_dir: Path, full_name: str, n: int) -> None:
    """Grava `n` releases sintéticas na janela (+ 1 pré-release e 1 fora da janela)."""
    owner, repo = full_name.split("/", 1)
    tags = [f"v{i}" for i in range(n)] + ["v-pre", "v-antiga"]
    datas = [datetime(2025, 11, 1 + i % 28, 12, 0, 0) for i in range(n)]
    datas += [datetime(2026, 1, 1), datetime(2024, 1, 1)]
    pre = [False] * n + [True, False]
    tabela = pa.table(
        {
            "tag_name": tags,
            "published_at": pa.array(datas, type=pa.timestamp("us", tz="UTC")),
            "prerelease": pre,
        }
    )
    releases_dir.mkdir(parents=True, exist_ok=True)
    pq.write_table(tabela, releases_dir / f"{owner}__{repo}.parquet")


@pytest.fixture
def cenario(tmp_path: Path, dummy_repos_df: pd.DataFrame):
    """Dataset sintético; sorteio original feito quando só metade dos repos tinha releases."""
    dataset = tmp_path / "repos.parquet"
    dummy_repos_df.to_parquet(dataset)
    sample = sortear_repositorios(dataset, n=NUM_REPOS, seed=DEFAULT_SEED)
    nomes = sample["full_name"].tolist()
    releases_dir = tmp_path / "releases"
    for i, nome in enumerate(nomes):
        _gravar_releases(releases_dir, nome, n=3 + 2 * i)  # alguns com <5, outros com >5
    saida = tmp_path / "amostra"
    parcial = tmp_path / "releases_parcial"
    for nome in nomes[: NUM_REPOS // 2]:
        owner, repo = nome.split("/", 1)
        shutil.copy(releases_dir / f"{owner}__{repo}.parquet", _mkdir(parcial))
    executar_sorteio_amostra_ouro(
        dataset_path=dataset,
        releases_dir=parcial,
        output_dir=saida,
        num_repos=NUM_REPOS,
        **JANELA,
    )
    return {"dataset": dataset, "releases": releases_dir, "saida": saida, "nomes": nomes, "tmp": tmp_path}


def _mkdir(p: Path) -> Path:
    p.mkdir(parents=True, exist_ok=True)
    return p


def _csv(saida: Path, nome: str = "amostra_ouro_template.csv") -> pd.DataFrame:
    return pd.read_csv(saida / nome, dtype=str, keep_default_na=False, encoding="utf-8-sig")


def _info(saida: Path) -> dict:
    return json.loads((saida / "sorteio_info.json").read_text(encoding="utf-8"))


def test_cenario_inicial_tem_placeholders_nos_repos_sem_releases(cenario):
    df = _csv(cenario["saida"])
    assert (df["tag_name"] == TAG_PLACEHOLDER).sum() == (NUM_REPOS // 2) * 5
    assert [bool(r["releases_sorteadas"]) for r in _info(cenario["saida"])["repositorios"]] == (
        [True] * (NUM_REPOS // 2) + [False] * (NUM_REPOS // 2)
    )


def test_completar_sorteio_pendente_preenche_so_os_pendentes(cenario):
    saida = cenario["saida"]
    antes_csv = _csv(saida)
    antes_info = _info(saida)

    res = completar_sorteio_pendente(cenario["dataset"], cenario["releases"], saida)

    assert res["alterado"] is True
    assert res["pendentes_completados"] == list(range(NUM_REPOS // 2 + 1, NUM_REPOS + 1))
    depois = _csv(saida)
    assert len(depois) == sum(len(r["releases_sorteadas"]) for r in _info(saida)["repositorios"])
    assert (depois["tag_name"] != TAG_PLACEHOLDER).all()
    assert (depois["published_at"] != "").all()
    assert depois["release_url"].str.contains("/releases/tag/").all()
    # linhas já reais: idênticas em conteúdo
    reais = antes_csv["tag_name"] != TAG_PLACEHOLDER
    pd.testing.assert_frame_equal(antes_csv[reais].reset_index(drop=True), depois[reais].reset_index(drop=True))
    # as 4 planilhas ficam idênticas entre si (em branco, cegas)
    for av in ("marcela", "gabriel", "joao"):
        pd.testing.assert_frame_equal(depois, _csv(saida, f"amostra_ouro_{av}.csv"))
        assert (saida / f"amostra_ouro_{av}.xlsx").is_file()
    for col in ("tipo_projeto", "entrega_real", "release_corretiva", "notas_justificativa"):
        assert (depois[col] == "").all()
    # info: sorteios anteriores intactos; pendentes preenchidos; registro da conclusão
    info = _info(saida)
    for antigo, novo in zip(antes_info["repositorios"], info["repositorios"]):
        if antigo["releases_sorteadas"]:
            assert novo == antigo
        else:
            assert 0 < len(novo["releases_sorteadas"]) <= 5
    assert info["timestamp_utc"] == antes_info["timestamp_utc"]
    assert info["seed"] == antes_info["seed"] == DEFAULT_SEED
    conclusao = info["conclusao_sorteio_pendente"]
    assert conclusao["repos_completados_idx"] == res["pendentes_completados"]
    assert conclusao["timestamp_utc"] and "random_state" in conclusao["metodo"]
    assert "pendente" in conclusao["nota"]


def test_completar_sorteio_pendente_e_igual_a_um_sorteio_completo_do_zero(cenario):
    completar_sorteio_pendente(cenario["dataset"], cenario["releases"], cenario["saida"])
    do_zero = cenario["tmp"] / "do_zero"
    executar_sorteio_amostra_ouro(
        dataset_path=cenario["dataset"],
        releases_dir=cenario["releases"],
        output_dir=do_zero,
        num_repos=NUM_REPOS,
        **JANELA,
    )

    pd.testing.assert_frame_equal(_csv(cenario["saida"]), _csv(do_zero))
    sorteadas = lambda d: [r["releases_sorteadas"] for r in _info(d)["repositorios"]]  # noqa: E731
    assert sorteadas(cenario["saida"]) == sorteadas(do_zero)


def test_completar_sorteio_pendente_e_idempotente(cenario):
    saida = cenario["saida"]
    completar_sorteio_pendente(cenario["dataset"], cenario["releases"], saida)
    arquivos = sorted(saida.iterdir())
    snapshot = {p.name: p.read_bytes() for p in arquivos}

    res = completar_sorteio_pendente(cenario["dataset"], cenario["releases"], saida)

    assert res["alterado"] is False and res["pendentes_completados"] == []
    assert {p.name: p.read_bytes() for p in sorted(saida.iterdir())} == snapshot


def test_completar_sorteio_pendente_sem_releases_persistidas_nao_grava_nada(cenario):
    saida = cenario["saida"]
    snapshot = {p.name: p.read_bytes() for p in saida.iterdir()}
    parcial = cenario["tmp"] / "releases_parcial"  # só a primeira metade dos repositórios

    with pytest.raises(SorteioInconsistenteError, match="sem releases persistidas"):
        completar_sorteio_pendente(cenario["dataset"], parcial, saida)

    assert {p.name: p.read_bytes() for p in saida.iterdir()} == snapshot


def test_completar_sorteio_pendente_recusa_releases_que_nao_reproduzem_o_sorteio_original(cenario):
    saida = cenario["saida"]
    snapshot = {p.name: p.read_bytes() for p in saida.iterdir()}
    primeiro = cenario["nomes"][0]
    _gravar_releases(cenario["releases"], primeiro, n=30)  # mudou o conjunto elegível

    with pytest.raises(SorteioInconsistenteError, match="não reproduzem o sorteio já registrado"):
        completar_sorteio_pendente(cenario["dataset"], cenario["releases"], saida)

    assert {p.name: p.read_bytes() for p in saida.iterdir()} == snapshot


def test_completar_sorteio_pendente_recusa_conjunto_de_repositorios_diferente(cenario):
    info = _info(cenario["saida"])
    info["repositorios"][0]["full_name"] = "outro/repo"
    (cenario["saida"] / "sorteio_info.json").write_text(json.dumps(info), encoding="utf-8")

    with pytest.raises(SorteioInconsistenteError, match="conjunto de repositórios"):
        completar_sorteio_pendente(cenario["dataset"], cenario["releases"], cenario["saida"])


def test_completar_sorteio_pendente_preserva_fim_de_linha_crlf(cenario):
    saida = cenario["saida"]
    for arq in saida.glob("*.csv"):
        arq.write_bytes(arq.read_bytes().replace(b"\r\n", b"\n").replace(b"\n", b"\r\n"))
    info = saida / "sorteio_info.json"
    info.write_bytes(info.read_bytes().replace(b"\r\n", b"\n").replace(b"\n", b"\r\n"))

    completar_sorteio_pendente(cenario["dataset"], cenario["releases"], saida)

    for arq in [*saida.glob("*.csv"), info]:
        dados = arq.read_bytes()
        assert dados.count(b"\n") == dados.count(b"\r\n"), arq.name
    assert _info(saida)["conclusao_sorteio_pendente"]


def test_amostra_ouro_versionada_nao_tem_placeholders():
    """O sorteio commitado está completo: 60 repositórios, 300 linhas reais."""
    base = Path(__file__).resolve().parent.parent / "dados" / "amostra_ouro"
    if not (base / "sorteio_info.json").is_file():
        pytest.skip("amostra-ouro ausente")
    info = json.loads((base / "sorteio_info.json").read_text(encoding="utf-8"))
    assert all(r["releases_sorteadas"] for r in info["repositorios"])
    for nome in ("template", "marcela", "gabriel", "joao"):
        df = pd.read_csv(base / f"amostra_ouro_{nome}.csv", dtype=str, keep_default_na=False, encoding="utf-8-sig")
        assert len(df) == 300 and (df["tag_name"] != TAG_PLACEHOLDER).all()
        assert (df["published_at"] != "").all()
